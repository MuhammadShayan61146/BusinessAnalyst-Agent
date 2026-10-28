"""Functions for creating Word documents."""
from docx import Document
from config.settings import OUTPUT_DIR

def save_doc(filename, title, sections, output_dir=None):
    """Create a Word document from a title and heading/line pairs.

    output_dir defaults to the legacy, fixed OUTPUT_DIR for backward
    compatibility with any existing caller that doesn't pass it
    (Phase 3B.2: graph nodes now pass the run-specific directory
    explicitly; this default exists only so nothing that already calls
    save_doc(filename, title, sections) breaks).
    """
    doc = Document()
    doc.add_heading(title, 0)
    for heading, lines in sections:
        doc.add_heading(heading, 1)
        for line in lines:
            doc.add_paragraph(str(line), style="List Bullet")
    path = f"{output_dir or OUTPUT_DIR}/{filename}"
    doc.save(path)
    return path
